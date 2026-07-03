"""Loaders: .txt / .md / .docx / raw string -> DocumentModel.

Docx loading extracts the page/typography metadata the docx-only rules need
(fonts, sizes, page geometry, endnote presence). Plain text just splits into
paragraphs on blank lines or newlines.
"""

from __future__ import annotations

import re
from pathlib import Path

from app.document import DocumentModel, Paragraph

_EMU_PER_INCH = 914_400  # python-docx lengths are EMU-backed


def load_text(text: str, source_format: str = "raw") -> DocumentModel:
    """Split raw text into paragraphs. Blank-line separation preferred;
    falls back to single newlines if the text has no blank lines."""
    if "\n\n" in text:
        blocks = re.split(r"\n\s*\n", text)
    else:
        blocks = text.splitlines()
    paragraphs = []
    idx = 0
    for block in blocks:
        # Keep leading whitespace (indentation matters for sub-listings)
        # but drop trailing whitespace and skip empty blocks.
        block = block.rstrip()
        if not block.strip():
            continue
        # Multi-line blocks: keep the block as one paragraph, joined.
        text_joined = "\n".join(line.rstrip() for line in block.splitlines())
        paragraphs.append(Paragraph(index=idx, text=text_joined))
        idx += 1
    return DocumentModel(paragraphs=paragraphs, source_format=source_format)


def _load_docx(path: Path) -> DocumentModel:
    import docx  # imported lazily so txt-only usage needs no python-docx

    d = docx.Document(str(path))
    paragraphs: list[Paragraph] = []
    fonts: set[str] = set()
    sizes: set[float] = set()

    for i, p in enumerate(d.paragraphs):
        if not p.text.strip():
            continue
        style_name = (p.style.name or "") if p.style else ""
        is_heading = style_name.lower().startswith("heading")
        runs_with_text = [r for r in p.runs if r.text.strip()]
        is_bold = bool(runs_with_text) and all(r.bold for r in runs_with_text)
        for r in runs_with_text:
            if r.font.name:
                fonts.add(r.font.name)
            if r.font.size is not None:
                sizes.add(r.font.size.pt)
        paragraphs.append(Paragraph(
            index=len(paragraphs), text=p.text,
            is_heading=is_heading, is_bold=is_bold,
        ))

    # Fall back to the document-default font if runs don't set one.
    try:
        normal = d.styles["Normal"].font
        if normal.name:
            fonts.add(normal.name)
        if normal.size is not None:
            sizes.add(normal.size.pt)
    except KeyError:
        pass

    sect = d.sections[0]
    metadata = {
        "fonts": fonts,
        "font_sizes": sizes,
        "page_width_in": sect.page_width / _EMU_PER_INCH if sect.page_width else None,
        "page_height_in": sect.page_height / _EMU_PER_INCH if sect.page_height else None,
        "margins_in": {
            "top": sect.top_margin / _EMU_PER_INCH,
            "bottom": sect.bottom_margin / _EMU_PER_INCH,
            "left": sect.left_margin / _EMU_PER_INCH,
            "right": sect.right_margin / _EMU_PER_INCH,
        },
        # Endnotes live in word/endnotes.xml inside the docx zip.
        "has_endnotes": _docx_part_has_notes(path, "endnotes"),
        "has_footnotes": _docx_part_has_notes(path, "footnotes"),
    }
    return DocumentModel(paragraphs=paragraphs, source_path=str(path),
                         source_format="docx", metadata=metadata)


def _docx_part_has_notes(path: Path, kind: str) -> bool:
    """True if word/{kind}.xml exists and contains real notes (beyond the
    built-in separator placeholders every docx carries)."""
    import zipfile

    try:
        with zipfile.ZipFile(path) as z:
            name = f"word/{kind}.xml"
            if name not in z.namelist():
                return False
            xml = z.read(name).decode("utf-8", errors="ignore")
            # Real notes have ids >= 1; separators use ids 0 and -1.
            return bool(re.search(r'w:id="[1-9]\d*"', xml))
    except (zipfile.BadZipFile, OSError):
        return False


def load_document(path: str | Path) -> DocumentModel:
    """Load a document by file extension (.txt, .md, .docx)."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _load_docx(path)
    if suffix in {".txt", ".md", ""}:
        doc = load_text(path.read_text(encoding="utf-8"),
                        source_format=suffix.lstrip(".") or "txt")
        doc.source_path = str(path)
        return doc
    raise ValueError(f"Unsupported file type: {suffix} "
                     "(supported: .txt, .md, .docx)")
